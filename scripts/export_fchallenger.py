#!/usr/bin/env python3.10
"""Exporta a planilha F CHALLENGER para o JSON de importação.

Uso:
    python3.10 scripts/export_fchallenger.py [xlsx] [saida.json]

Padrões:
    xlsx   = ~/Downloads/F CHALLENGER R25.xlsx
    saida  = /tmp/opencode/fchallenger.json   (NÃO vai para o repositório:
             contém telefones dos atletas — dado pessoal)

O JSON contém:
  - athletes: atletas canônicos (união ATLETAS + RANKING + nomes de jogos),
    com nome de exibição, telefone e nível da planilha;
  - matches : os 84 jogos com as correções evidentes já aplicadas
    (linhas 19 e 76 com ano errado; linha 24 com sets espelhados);
  - official: o ranking oficial da aba RANKING (só para conferência —
    o app recalcula tudo pelas regras).

Correspondência com a análise /tmp/opencode/analyze_fchallenger.py.
"""
import datetime
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl

DEFAULT_XLSX = os.path.expanduser("~/Downloads/F CHALLENGER R25.xlsx")
DEFAULT_OUT = "/tmp/opencode/fchallenger.json"


def norm(s):
    if s is None:
        return ""
    s = str(s).strip().upper()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# grafias diferentes entre abas para a MESMA pessoa (normalizadas)
ALIAS = {
    "LUIZ GLAESER": "LUIS GLAESER", "LUIS H GLAESER": "LUIS GLAESER",
    "LUIZ H GLAESER": "LUIS GLAESER", "LUIZ H. GLAESER": "LUIS GLAESER",
    "FAGNER LUIZ DE OLIVEIRA": "FAGNER OLIVEIRA",
    "LUIS ARTUR ROSSA": "LUIS ROSSA", "LUIS ARTHUR ROSSA": "LUIS ROSSA",
    "LUIS ARHUR ROSSA": "LUIS ROSSA", "LUIZ ARTHUR ROSSA": "LUIS ROSSA",
    "GUILHERME ROMBALDI NUNES": "GUILHERME R NUNES",
    "ARALDO NETO": "ARALDO",
    "WILLIAN GIL": "WILLIAM GIL",
    "JESSICA BONFIM": "JESSICA BOMFIM",
    "KASSIANO LOUZADA": "KASSIANO",
    "VITOR FURMANN": "VITOR",
}

# apelidos usados apenas no REGISTRO DE JOGOS (evidência: batem com o oficial)
LOG_ALIAS = {
    "RODRIGO": "RODRIGO VIEIRA",
    "DIEGO": "DIEGO DUARTE",
    "GUILHERME": "GUILHERME R NUNES",
    "MAICON": "MAICON BOSCAINI",
    "TIAGO": "TIAGO BARBOSA",
}

# Nome de exibição preferido por canônico (chaves completas da planilha
# que o norm() encurta — o norm deles é diferente, por isso não é scan).
DISPLAY_PREFERRED = {
    "VITOR": "VITOR FURMANN",
    "ARALDO": "ARALDO NETO",
    "KASSIANO": "KASSIANO LOUZADA",
    "GUILHERME R NUNES": "GUILHERME ROMBALDI NUNES",
    "LUIS ROSSA": "LUIS ARTHUR ROSSA",
    "LUIS GLAESER": "LUIS H GLAESER",
    "FAGNER OLIVEIRA": "FAGNER LUIZ DE OLIVEIRA",
}

PARTICLES = {"DE", "DA", "DO", "DAS", "DOS", "E"}


def titlecase(upper):
    words = []
    for i, w in enumerate(upper.split(" ")):
        if i > 0 and w in PARTICLES:
            words.append(w.lower())
        else:
            words.append(w.capitalize() if w else w)
    return " ".join(words)


def canon(name):
    n = norm(name)
    return ALIAS.get(n, n)


def log_canon(name):
    return LOG_ALIAS.get(canon(name), canon(name))


def parse_sets(*cells):
    sets = []
    for c in cells:
        if c is None or str(c).strip() == "":
            continue
        ints = [int(x) for x in re.findall(r"\d+", str(c))]
        if len(ints) >= 2:
            sets.append((ints[0], ints[1]))
    return sets


def points_normal(sa, sb):
    if sa == 2 and sb == 0:
        return 300, 50
    if sa == 2 and sb == 1:
        return 200, 100
    if sa == 1 and sb == 2:
        return 100, 200
    if sa == 0 and sb == 2:
        return 50, 300
    return None, None


GS_MAP = {0: 75, 1: 150, 2: 300, 3: 450}


def points_gs(sa, sb):
    if max(sa, sb) != 3 or min(sa, sb) > 2:
        return None, None
    return GS_MAP[sa], GS_MAP[sb]


def parse_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return datetime.date(v.year, v.month, v.day)
    if isinstance(v, str):
        m = re.match(r"\s*(\d{1,2})\.(\d{1,2})\.(\d{4})", v)
        if m:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    out = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    problems = []

    # ---------- ATLETAS (A=nº B=ATLETA C=NÍVEL D=TELEFONE; dados na linha 6) ----------
    athletes = {}  # canon -> {"variants": [...], "nivel", "telefone"}
    ws = wb["ATLETAS"]
    for row in ws.iter_rows(min_row=6):
        nome = row[1].value
        if nome is None or str(nome).strip() == "":
            continue
        key = canon(nome)
        rec = athletes.setdefault(
            key, {"variants": [], "nivel": None, "telefone": None}
        )
        raw = str(nome).strip()
        if norm(raw) == key:  # variante limpa (sem "typos" de alias)
            rec["variants"].append(raw)
        rec["nivel"] = row[2].value or rec["nivel"]
        tel = row[3].value if len(row) > 3 else None
        if tel:
            rec["telefone"] = str(tel).strip()

    # ---------- RANKING oficial ----------
    official, order = {}, []
    ws = wb["RANKING"]
    for row in ws.iter_rows(min_row=7, max_row=79):
        pos, nome, total, jgs = row[2].value, row[3].value, row[6].value, row[11].value
        if nome is None or not isinstance(total, (int, float)):
            continue
        key = canon(nome)
        official[key] = {
            "pos": pos,
            "total": int(total),
            "jgs": int(jgs),
            "nome": str(nome).strip(),
        }
        order.append(key)
        rec = athletes.setdefault(
            key, {"variants": [], "nivel": None, "telefone": None}
        )
        raw = str(nome).strip()
        if norm(raw) == key:
            rec["variants"].append(raw)

    # ---------- REGISTRO DE JOGOS ----------
    ws = wb["REGISTRO DE JOGOS"]
    games = []
    for row in ws.iter_rows(min_row=8):
        v = [c.value for c in row]
        # A=nº B=DATA C=HORÁRIO D=LOCAL E=QUADRA F=JOG A G=JOG B
        # H/I/J sets K=VIT L/M pts N=OBS
        data = v[1]
        a, b = v[5], v[6]
        h, i, j = v[7], v[8], v[9]
        obs = v[13] if len(v) > 13 else None
        if a is None or b is None:
            continue
        dt = parse_date(data)
        sets = parse_sets(h, i, j)
        # Correções de digitação evidentes (ver análise):
        #  linha 19: data 06.01.2025 -> 2026 (jogo entre jogos de jan/2026)
        #  linha 76: data 07.09.2025 -> 2026 (junto a outros de set/2026)
        #  linha 24: sets lançados espelhados; K (ALESSANDRO) e os pontos
        #            oficiais concordam com vitória de ALESSANDRO -> espelhar
        if row[0].row == 19:
            dt = datetime.date(2026, 1, 6)
        if row[0].row == 76:
            dt = datetime.date(2026, 9, 7)
        if row[0].row == 24:
            sets = [(0, 6), (2, 6)]
        if dt is None:
            problems.append(f"linha {row[0].row}: sem data")
            continue
        if dt < datetime.date(2025, 11, 1):
            problems.append(
                f"linha {row[0].row}: data {dt} anterior ao início (nov/2025)"
            )
        a_key, b_key = log_canon(a), log_canon(b)
        sa = sum(1 for x, y in sets if x > y)
        sb = sum(1 for x, y in sets if y > x)
        is_gs = bool(obs and "GRAND SLAM" in str(obs).upper())
        pa, pb = points_gs(sa, sb) if is_gs else points_normal(sa, sb)
        if pa is None:
            problems.append(
                f"linha {row[0].row}: placar sem vencedor válido "
                f"({a_key} x {b_key}, sets={sets}, gs={is_gs})"
            )
        for k in (a_key, b_key):
            athletes.setdefault(
                k, {"variants": [], "nivel": None, "telefone": None}
            )
        games.append(
            {
                "row": row[0].row,
                "date": dt.isoformat(),
                "a_key": a_key,
                "b_key": b_key,
                "sets": [[x, y] for x, y in sets],
                "gs": is_gs,
                "calc": [pa, pb],
            }
        )

    # ---------- Nome de exibição ----------
    # base = canônico (title case); variante mais longa de ATLETAS/RANKING
    # vence; PREFERRED cobre as chaves completas encurtadas pelo norm().
    display = {}
    for key, rec in athletes.items():
        cands = [DISPLAY_PREFERRED.get(key, key)]
        cands.extend(rec["variants"])
        best = max(cands, key=len)  # empate -> primeiro (o preferred)
        display[key] = titlecase(best)

    # ---------- Conferência: recálculo vs RANKING oficial ----------
    tot, cnt = defaultdict(int), defaultdict(int)
    for g in games:
        pa, pb = g["calc"]
        if pa is None:
            continue
        tot[g["a_key"]] += pa
        cnt[g["a_key"]] += 1
        tot[g["b_key"]] += pb
        cnt[g["b_key"]] += 1

    divergences = []
    for k in order:
        off = official[k]
        if tot.get(k, 0) != off["total"] or cnt.get(k, 0) != off["jgs"]:
            divergences.append(
                f"{display.get(k, k)}: oficial={off['total']}/{off['jgs']}g "
                f"recalc={tot.get(k, 0)}/{cnt.get(k, 0)}g"
            )

    out_data = {
        "source": os.path.basename(xlsx),
        "exported_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "athletes": [
            {
                "key": k,
                "name": display[k],
                "phone": rec["telefone"],
                "ref_level": str(rec["nivel"]).strip() if rec["nivel"] else None,
                "in_games": k in tot,
            }
            for k, rec in sorted(athletes.items())
        ],
        "matches": [
            {
                "date": g["date"],
                "a_key": g["a_key"],
                "b_key": g["b_key"],
                "a": display.get(g["a_key"], g["a_key"]),
                "b": display.get(g["b_key"], g["b_key"]),
                "sets": g["sets"],
                "format": "grand_slam" if g["gs"] else "normal",
            }
            for g in games
        ],
        "official": [
            {
                "key": k,
                "name": display.get(k, k),
                "pos": official[k]["pos"],
                "total": official[k]["total"],
                "jgs": official[k]["jgs"],
            }
            for k in order
        ],
    }

    with open(out, "w", encoding="utf-8") as f:
        json.dump(out_data, f, ensure_ascii=False, indent=1)

    in_games = sum(1 for a in out_data["athletes"] if a["in_games"])
    print(f"OK -> {out}")
    print(
        f"atletas={len(out_data['athletes'])} "
        f"(com jogos: {in_games}) jogos={len(out_data['matches'])} "
        f"oficial={len(out_data['official'])} "
        f"GS={sum(1 for m in out_data['matches'] if m['format'] == 'grand_slam')}"
    )
    if divergences:
        print("divergências recalc vs oficial (esperadas: RODRIGO MEINCKE):")
        for d in divergences:
            print(f"  - {d}")
    if problems:
        print("PROBLEMAS (não exportados ou a corrigir):")
        for p in problems:
            print(f"  ! {p}")
        sys.exit(1)


if __name__ == "__main__":
    main()
